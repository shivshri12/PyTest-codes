import openpyxl
import openpyxl.styles
from openpyxl.styles import PatternFill

def noofrows(file,sheet):
    workbook = openpyxl.load_workbook(file)
    ws = workbook[sheet]
    return ws.max_row

def noofcolumns(file,sheet):
    workbook = openpyxl.load_workbook(file)
    ws = workbook[sheet]
    return ws.max_column

def readdata(file,sheet,row,column):
    workbook = openpyxl.load_workbook(file)
    ws = workbook[sheet]
    return ws.cell(row,column).value

def writedata(file,sheet,row,column,data):
    workbook = openpyxl.load_workbook(file)
    ws = workbook[sheet]
    ws.cell(row,column).value=data
    workbook.save(file)

def fillGreencolor(file,sheet,row,column):
    workbook = openpyxl.load_workbook(file)
    ws = workbook[sheet]
    fill=PatternFill(start_color='00FF00',end_color='00FF00',fill_type='solid')
    ws.cell(row,column).fill=fill
    workbook.save(file)

def fillRedcolor(file,sheet,row,column):
    workbook = openpyxl.load_workbook(file)
    ws = workbook[sheet]
    fill=PatternFill(start_color='FF0000',end_color='FF0000',fill_type='solid')
    ws.cell(row,column).fill=fill
    workbook.save(file)

